from ListaAlumnos import gpoPB
from openpyxl import Workbook
from openpyxl import load_workbook

try:
    libro = load_workbook("Gpo073EJ25.xlsx")
except FileNotFoundError:
    print("El archivo no existe, adiós!")
    exit()

if len(libro.sheetnames)>1:
    libro.active = 1
    hoja = libro.active
else:
    hoja = libro.active
print ("La página activa es:",hoja.title)

celda = hoja["A1"]
hoja["A1"] = "Num. lista"
print ("Contenido de A1:",celda.value)
hoja["B1"] = "Matricula"
hoja["C1"] = "Alumno"
hoja["D1"] = "Semestre"
hoja["E1"] = "Carrera"
print ("Los encabezados de las celdas son:")
print (hoja["A1"].value,end = "\t")
print (hoja["B1"].value, end = "\t")
print (hoja["C1"].value, end = "\t")
print (hoja["D1"].value, end = "\t")
print (hoja["E1"].value)
input()
count = 2
for i in range(len(gpoPB)):
    hoja.cell(count,1,i+1) #i+2
    hoja.cell(count,2,gpoPB[i]['mat'])
    hoja.cell(count,3,gpoPB[i]['nombre'])
    hoja.cell(count,4,gpoPB[i]['sem'])
    hoja.cell(count,5,gpoPB[i]['carrera'])
    count += 1

libro.save("Gpo073EJ25.xlsx")
