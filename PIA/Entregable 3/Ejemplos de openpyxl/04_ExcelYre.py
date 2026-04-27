from openpyxl import Workbook
from openpyxl import load_workbook
import re

date = re.compile (r"\w{3}\s{1,2}\d{1,2}") #May 18     May  5
hour = re.compile (r"\d{2}\:\d{2}:\d{2}") # 13:30:55
numlinea = 1

#Creamos el Excel
libro = Workbook()
hoja = libro.active

fo = open ("auth_log.txt")
for line in fo:
    d = date.search(line)
    h = hour.search(line)
    hoja.cell(numlinea,1, d.group()) #NUMLINEA indice fila, # indico la columna
    hoja.cell(numlinea,2, h.group())
    numlinea += 1
fo.close()
libro.save("Mi ejercicio3.xls")
